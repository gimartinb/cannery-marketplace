from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

out = Path('/home/ubuntu/cannery-marketplace/Vendor-Bulk-Import-Template.xlsx')
wb = Workbook()
ws = wb.active
ws.title = 'Vendor Import'
headers = [
    'name*','slug','category*','short_bio*','owner_bio','contact_email*',
    'website_url','instagram_url','facebook_url','tiktok_url','cover_image_url',
    'gallery_image_urls','featured','display_order','active'
]
ws.append(headers)
ws.append([
    'Example Maker Studio','example-maker-studio','Ceramics',
    'Handmade goods created locally in small batches.',
    'Tell shoppers about the person behind the work.','maker@example.com',
    'https://example.com','https://www.instagram.com/examplemaker','', '', '',
    'https://example.com/photo-1.jpg | https://example.com/photo-2.jpg',
    'No',10,'Yes'
])
navy = '20241D'; gold = 'A87816'; cream = 'F5F0E6'; white = 'FFFDF8'; red = 'B5563B'
for cell in ws[1]:
    cell.fill = PatternFill('solid', fgColor=navy)
    cell.font = Font(color=white, bold=True)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = Border(bottom=Side(style='thin', color=gold))
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:O2'
widths = [28,28,22,46,46,30,34,36,36,36,36,50,14,16,12]
for i, width in enumerate(widths, start=1):
    ws.column_dimensions[chr(64+i)].width = width
ws.row_dimensions[1].height = 34
for row in ws.iter_rows(min_row=2, max_row=250, min_col=1, max_col=len(headers)):
    for cell in row:
        cell.alignment = Alignment(vertical='top', wrap_text=True)
        if cell.row % 2 == 0:
            cell.fill = PatternFill('solid', fgColor=cream)
yes_no = DataValidation(type='list', formula1='"Yes,No"', allow_blank=True)
ws.add_data_validation(yes_no)
yes_no.add('M2:M250'); yes_no.add('O2:O250')
order = DataValidation(type='whole', operator='between', formula1='0', formula2='9999', allow_blank=True)
ws.add_data_validation(order); order.add('N2:N250')

instructions = wb.create_sheet('Instructions')
instructions['A1'] = 'The Cannery Marketplace — One-Time Vendor Import'
instructions['A1'].font = Font(size=18, bold=True, color=navy)
instructions['A3'] = 'How to use this file'
instructions['A3'].font = Font(size=13, bold=True, color=gold)
steps = [
    '1. Keep the header row unchanged. Fields marked with * are required.',
    '2. Add one vendor per row. Delete the example row before importing real vendors.',
    '3. Social links must point to the matching service (Instagram, Facebook, or TikTok).',
    '4. Separate multiple gallery image URLs with a vertical bar: URL 1 | URL 2. Maximum five gallery images.',
    '5. Save as .xlsx, upload in Admin > Vendor import, review all errors, then confirm the import.',
    '6. Imported vendors are approved by the administrator and can appear publicly when Active is Yes.'
]
for idx, text in enumerate(steps, start=4):
    instructions.cell(idx, 1, text)
instructions['A12'] = 'Security and privacy'
instructions['A12'].font = Font(size=13, bold=True, color=red)
instructions['A13'] = 'Do not include passwords, payment information, government IDs, or private notes in this file.'
instructions['A15'] = 'Allowed social domains'
instructions['A15'].font = Font(size=13, bold=True, color=gold)
instructions['A16'] = 'Instagram: instagram.com  |  Facebook: facebook.com or fb.com  |  TikTok: tiktok.com'
instructions.column_dimensions['A'].width = 118
for row in range(1, 18):
    instructions.cell(row,1).alignment = Alignment(wrap_text=True, vertical='top')
instructions.sheet_view.showGridLines = False
wb.save(out)
print(out)
